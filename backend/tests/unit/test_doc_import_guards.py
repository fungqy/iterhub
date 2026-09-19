"""文档导入的「护栏」:把 2026-09-17 那轮排查修掉的四处非判定类问题钉住。

与 test_doc_bug_validate.py 的分工:那个文件守的是**判定口径**(行级必填、故事号解析、
strict/lenient 的同源性);这里守的是**外圈约定** —— 上传格式白名单、sheet 名在文本与
截图两侧的一致性、openpyxl 补丁的生命周期、以及「重活不许写进 async def」。

为什么这四条值得单独钉:它们都属于「删掉/改回去都不会报错,只会安静地坏」的那类约定 ——
比如有人把 .xls 顺手加回白名单、或把路由改回 async def,单测全绿而线上出问题。
"""

from __future__ import annotations

import inspect
from pathlib import Path

import pytest
from fastapi import HTTPException, UploadFile
from openpyxl import Workbook
from openpyxl.styles.fills import Fill

from api.routes import data_import as di


def _upload(filename: str) -> UploadFile:
    import io

    return UploadFile(file=io.BytesIO(b""), filename=filename)


# ══ 上传格式白名单 ══════════════════════════════════════════════
class TestRequireXlsx:
    def test_accepts_xlsx_case_insensitive(self):
        """`.XLSX` 是合法后缀(浏览器 accept 本来就忽略大小写),不能拒。"""
        di._require_xlsx(_upload("现场问题.xlsx"))
        di._require_xlsx(_upload("A.XLSX"))

    @pytest.mark.parametrize("name", ["x.xls", "x", "x.xlsx.bak", "x.xlsm", ""])
    def test_rejects_everything_else_with_400(self, name):
        """⚠ 老版 .xls 必须在这一步被 400 拦下:openpyxl 读不了 OLE2,放过去就是
        一个必然失败的 500「处理文件失败」——而界面此前还在宣传支持 .xls。"""
        with pytest.raises(HTTPException) as ei:
            di._require_xlsx(_upload(name))
        assert ei.value.status_code == 400
        assert ".xlsx" in ei.value.detail


# ══ sheet 名:文本与截图必须读同一张 ════════════════════════════
class TestSheetNameThreaded:
    def test_returns_the_sheet_actually_read(self, tmp_path: Path):
        """表名不符时退用活动表 —— 那个**实际用的名字**必须回传给截图侧。

        否则文本从一个 sheet 读、drawing 按另一个名字找 ⇒ 行照导、图一张不入,
        理由还写着「源文件无内嵌图片」。
        """
        wb = Workbook()
        ws = wb.active
        ws.title = "明细"
        ws.append(["自动编号", "问题详述", "sprint"])
        ws.append(["m0900001", "点不动", "botadp-Sprint-11"])
        path = tmp_path / "renamed.xlsx"
        wb.save(path)

        headers, rows, strikethrough, sheet_name = di.read_doc_bug_sheet(str(path))
        assert sheet_name == "明细"
        assert headers[:3] == ["自动编号", "问题详述", "sprint"]
        assert [r[0] for r in rows] == [2]
        assert strikethrough == 0

    def test_expected_sheet_name_is_used_when_present(self, tmp_path: Path):
        wb = Workbook()
        ws = wb.active
        ws.title = di.DOC_BUG_SHEET_NAME
        ws.append(["自动编号", "问题详述", "sprint"])
        path = tmp_path / "ok.xlsx"
        wb.save(path)

        assert di.read_doc_bug_sheet(str(path))[3] == di.DOC_BUG_SHEET_NAME


# ══ openpyxl 容错补丁的生命周期 ════════════════════════════════
class TestFillPatchIsProcessWide:
    def test_patch_is_permanent_not_scoped(self):
        """补丁必须是**进程级常驻**,不能再退回 load_workbook 前后的 patch/unpatch。

        那是改进程级类属性:并发导入时先结束的请求会把补丁撤掉,另一个还在解析途中,
        于是又回到 TypeError、报成随机的 500(路由改成线程池后并发是常态)。
        """
        assert Fill.__init__.__name__ == "_tolerant_fill_init"
        assert Fill.__init__ is di._tolerant_fill_init

    def test_malformed_fill_does_not_raise(self):
        """补丁的实际作用:构造时抛的 TypeError 被吞掉,字段回落成 None。

        真实触发路径:源表里有 `<fill>` 缺 patternType/fgColor,openpyxl 的 styles 读取器
        拿着残缺属性去构造 Fill。这里用一个意料之外的 kwarg 复现同一类 TypeError
        (原版会 `TypeError: unexpected keyword argument`,整份文件读不下去)。
        """
        fill = Fill(bogus_attribute=1)
        assert fill.patternType is None
        assert fill.fgColor is None and fill.bgColor is None


# ══ 重活不许写进 async def ═════════════════════════════════════
class TestRoutesStaySync:
    """⚠ 这两个接口全程是同步重活(落盘 / openpyxl 解析几百 MB / Pillow 逐张解码 /
    逐张写库)。写成 async def 时它们跑在事件循环线程上 —— 一次导入会把全服其它接口
    卡住数秒到数十秒。写成 def 由 FastAPI 丢线程池,是刻意的,别"顺手"改回去。
    """

    @pytest.mark.parametrize("handler", [di.upload_doc_bugs, di.validate_doc_bugs])
    def test_not_coroutine_function(self, handler):
        assert not inspect.iscoroutinefunction(handler), (
            f"{handler.__name__} 又被写回 async def 了:同步重活会独占事件循环"
        )


# ══ sprint_id 在端点上必须是必填 ═══════════════════════════════
class TestSprintIdIsRequired:
    """2026-09-19:导入范围由「整个项目」收窄到「所选迭代」,界面在第 2 步就让用户选。

    ⚠ 为什么宁可 422 也不给默认值:sprint_id 漏传时的行为恰恰是最坏的那种 ——
      后端会把**整个项目**的行都导进来,界面完全看不出异常,只有事后对比行数
      才发现多导了几个迭代(而库里那些行的 sprint_id 各自是对的,所以列出来也正常)。

    ⚠ 这里看的是 `Query(...)` 里的默认值,不是 Python 签名里的默认值:
      端点写的是 `sprint_id: str = Query(...)`,语法上"有默认值",语义上"必填" ——
      所以**不能**断言 `param.default is inspect.Parameter.empty`(那个永远为假)。

    ⚠ 也不能拿 `is Ellipsis` 判"必填":本环境(FastAPI 0.11x / pydantic v2)的
      `Query(...)` 把 default 存成 `PydanticUndefined`,不是 Ellipsis —— 第一版断言
      就是这么写错、红了一次。改用 FieldInfo.is_required(),由库自己回答"要不要传"。
    """

    @pytest.mark.parametrize("handler", [di.upload_doc_bugs, di.validate_doc_bugs])
    def test_sprint_id_has_no_default(self, handler):
        sig = inspect.signature(handler)
        assert "sprint_id" in sig.parameters, f"{handler.__name__} 缺 sprint_id 参数"
        default = sig.parameters["sprint_id"].default
        is_required = getattr(default, "is_required", None)
        assert callable(is_required) and is_required(), (
            f"{handler.__name__} 的 sprint_id 不是必填 —— 漏传会静默把范围放大到整个项目"
        )
